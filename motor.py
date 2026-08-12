import os
import shutil
import tempfile
import unicodedata

from openpyxl import load_workbook


TEMPLATE_EXCEL = (
    "Excel para alimentar o powerpoint - Preenchido.xlsx"
)


def encontrar_operadoras(arquivo_excel):

    wb = load_workbook(
        arquivo_excel,
        data_only=True
    )

    operadoras = set()

    operadoras_conhecidas = [
        "Amil Selecionada",
        "Amil Metal",
        "Bradesco",
        "SulAmérica",
        "Omint",
        "Hapvida",
        "Seguros Unimed",
    ]

    for nome_aba in [
        "Matriz FX Et. C COPART",
        "Matriz FX Et. S COPART",
    ]:

        if nome_aba not in wb.sheetnames:
            continue

        ws = wb[nome_aba]

        for row in ws.iter_rows():

            for cell in row:

                valor = str(
                    cell.value or ""
                ).strip()

                if not valor:
                    continue

                for operadora in operadoras_conhecidas:

                    if (
                        valor.casefold()
                        == operadora.casefold()
                    ):
                        operadoras.add(
                            operadora
                        )

    wb.close()

    ordem_preferencial = [
        "Amil Selecionada",
        "Amil Metal",
        "Bradesco",
        "Hapvida",
        "Omint",
        "SulAmérica",
        "Seguros Unimed",
    ]

    return [
        operadora
        for operadora in ordem_preferencial
        if operadora in operadoras
    ]
# =====================================================
# FUNÇÕES AUXILIARES DA ABA BASE
# =====================================================


def normalizar_base(valor):
    if valor is None:
        return ""

    texto = str(valor).strip().upper()

    texto = unicodedata.normalize(
        "NFKD",
        texto
    )

    texto = "".join(
        caractere
        for caractere in texto
        if not unicodedata.combining(
            caractere
        )
    )

    return " ".join(
        texto.split()
    )


def localizar_cabecalhos_base(ws):
    """
    Localiza a linha de cabeçalhos da aba Base
    sem depender de letras fixas de coluna.
    """

    cabecalhos = {}

    limite_linhas = min(
        ws.max_row,
        20
    )

    for linha in range(
        1,
        limite_linhas + 1
    ):
        valores_linha = [
            normalizar_base(
                ws.cell(
                    linha,
                    coluna
                ).value
            )
            for coluna in range(
                1,
                ws.max_column + 1
            )
        ]

        if (
            "NOME" in valores_linha
            and
            "FAIXA ETARIA" in valores_linha
        ):
            for coluna in range(
                1,
                ws.max_column + 1
            ):
                texto = normalizar_base(
                    ws.cell(
                        linha,
                        coluna
                    ).value
                )

                if not texto:
                    continue

                cabecalhos.setdefault(
                    texto,
                    []
                ).append(
                    coluna
                )

            return (
                linha,
                cabecalhos
            )

    raise ValueError(
        "Não foi possível localizar o "
        "cabeçalho da aba Base."
    )


def primeira_coluna_base(
    cabecalhos,
    nomes
):
    for nome in nomes:
        nome_normalizado = normalizar_base(
            nome
        )

        colunas = cabecalhos.get(
            nome_normalizado,
            []
        )

        if colunas:
            return colunas[0]

    return None


def ultima_coluna_base(
    cabecalhos,
    nomes
):
    for nome in nomes:
        nome_normalizado = normalizar_base(
            nome
        )

        colunas = cabecalhos.get(
            nome_normalizado,
            []
        )

        if colunas:
            return colunas[-1]

    return None

def gerar_excel_final(
    caminho_matriz,
    operadoras_em_ordem
):
    """
    Primeira integração da aplicação.

    Nesta versão:
    1. valida a matriz;
    2. valida as operadoras escolhidas;
    3. cria uma cópia independente do Excel modelo;
    4. grava a ordem das operadoras no arquivo;
    5. devolve o caminho para download.

    O preenchimento completo das abas será conectado
    nesta mesma função, sem arquivos 7A, 7B, 7C ou 7D.
    """

    if not os.path.exists(caminho_matriz):
        raise FileNotFoundError(
            "A matriz temporária não foi encontrada."
        )

    if not os.path.exists(TEMPLATE_EXCEL):
        raise FileNotFoundError(
            "O Excel modelo não foi encontrado no projeto: "
            f"{TEMPLATE_EXCEL}"
        )

    if not operadoras_em_ordem:
        raise ValueError(
            "Selecione pelo menos uma operadora."
        )

    if len(operadoras_em_ordem) > 4:
        raise ValueError(
            "Selecione no máximo quatro operadoras."
        )

    if (
        len(operadoras_em_ordem)
        != len(set(operadoras_em_ordem))
    ):
        raise ValueError(
            "Não é permitido repetir operadoras."
        )

    operadoras_encontradas = encontrar_operadoras(
        caminho_matriz
    )

    invalidas = [
        operadora
        for operadora in operadoras_em_ordem
        if operadora not in operadoras_encontradas
    ]

    if invalidas:
        raise ValueError(
            "As seguintes operadoras não foram "
            f"encontradas na matriz: {invalidas}"
        )

    pasta_saida = tempfile.mkdtemp(
        prefix="apresentador360_"
    )

    caminho_saida = os.path.join(
        pasta_saida,
        "Apresentador360_Excel_Final.xlsx"
    )

    shutil.copy2(
        TEMPLATE_EXCEL,
        caminho_saida
    )

    wb = load_workbook(
        caminho_saida,
        data_only=False
    )

    # Registra a seleção no arquivo para validar
    # a conexão Streamlit → motor → Excel.
    #
    # A aba Planos já existe no modelo.
    ws = wb["Planos"]

    # Limpa a área de controle.
    for linha in range(1, 10):
        ws.cell(
            linha,
            20
        ).value = None

        ws.cell(
            linha,
            21
        ).value = None

    ws.cell(
        1,
        20
    ).value = "ORDEM APRESENTADOR 360"

    for posicao, operadora in enumerate(
        operadoras_em_ordem,
        start=1
    ):
        ws.cell(
            posicao + 1,
            20
        ).value = posicao

        ws.cell(
            posicao + 1,
            21
        ).value = operadora

    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    except Exception:
        pass

    wb.save(caminho_saida)
    wb.close()

    return caminho_saida
